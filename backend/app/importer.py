"""Excel → SQLite import mantig'i.

CLI skript (scripts/import_excel.py) va admin API (routers/admin.py)
ikkalasi shu moduldan foydalanadi — parsing va tozalash bir joyda turadi.
"""

import json
import sqlite3
from io import BytesIO
from pathlib import Path

import openpyxl

from .normalize import clean, normalize

DB_PATH = Path(__file__).resolve().parents[1] / "data" / "tezaurus.db"

TYPE_PREFIX = {"allyuziv-nom": "an", "iqtibos": "iq"}

SCHEMA = """
CREATE TABLE IF NOT EXISTS entries (
    id                  TEXT PRIMARY KEY,
    type                TEXT NOT NULL,          -- 'allyuziv-nom' | 'iqtibos'
    unit                TEXT NOT NULL,
    unit_normalized     TEXT NOT NULL,
    pronunciations      TEXT NOT NULL,          -- JSON massiv
    context_text        TEXT NOT NULL,
    intertext_author    TEXT NOT NULL,
    intertext_work      TEXT NOT NULL,
    intertext_genre     TEXT NOT NULL,
    intertext_year      TEXT NOT NULL,
    intertext_publisher TEXT NOT NULL,
    intertext_page      TEXT NOT NULL,
    source_description  TEXT NOT NULL,
    source_author       TEXT NOT NULL,
    source_work         TEXT NOT NULL,
    source_genre        TEXT NOT NULL,
    source_publisher    TEXT NOT NULL,
    source_period       TEXT NOT NULL,
    recognition         TEXT NOT NULL,          -- 'yadro' | 'periferiya' | ''
    commentary          TEXT NOT NULL,
    semantic_field      TEXT NOT NULL,
    synonyms            TEXT NOT NULL,          -- JSON massiv
    hypernym            TEXT NOT NULL,
    hyponym             TEXT NOT NULL,
    note                TEXT NOT NULL,
    is_complete         INTEGER NOT NULL        -- asl manba/sharh to'ldirilganmi
);
"""


def parse_recognition(value: str) -> str:
    v = normalize(value)
    if v.startswith("yadro"):
        return "yadro"
    if v.startswith(("perefer", "perifer")):
        return "periferiya"
    return ""


def parse_semantic_field(value: str) -> str:
    v = clean(value)
    return v[:1].upper() + v[1:] if v else ""


def parse_pronunciations(value: str) -> list[str]:
    parts = [clean(p) for p in str(value or "").replace("\\", "/").split("//")]
    # ba'zan bitta / bilan ham ajratilgan
    result: list[str] = []
    for p in parts:
        result.extend(clean(x) for x in p.split("/") if clean(x))
    seen, uniq = set(), []
    for p in result:
        key = normalize(p)
        if key and key not in seen:
            seen.add(key)
            uniq.append(p)
    return uniq


def parse_synonyms(value: str) -> list[str]:
    """Vergul bilan ajratilgan qisqa ro'yxatni bo'lish; uzun jumlani butunligicha saqlash."""
    v = clean(value)
    if not v:
        return []
    parts = [clean(p) for p in v.split(",") if clean(p)]
    if len(parts) > 1 and all(len(p) <= 40 for p in parts):
        return parts
    return [v]


def rows_from_workbook(wb, type_: str) -> tuple[list[dict], dict]:
    """Ochilgan workbook'dan yozuvlarni o'qish.

    Qaytaradi: (entries, hint) — hint 3-ustundagi tur yorlig'i statistikasi
    (fayl noto'g'ri turda yuklanganini aniqlash uchun).
    """
    prefix = TYPE_PREFIX[type_]
    ws = wb.active
    entries: list[dict] = []
    hint = {"allyuziv-nom": 0, "iqtibos": 0}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cells = [clean(c) for c in row[:24]]
        cells += [""] * (24 - len(cells))
        (rid, unit, unit_label, pron, ctx_text,
         it_author, it_work, it_genre, it_year, it_pub, it_page,
         src_desc, src_author, src_work, src_genre, src_pub, src_period,
         recognition, commentary, sem_field, synonym, hypernym, hyponym, note) = cells
        if not unit:  # faqat ID qolgan texnik qatorlar
            continue
        label = normalize(unit_label)
        if label.startswith("allyuz"):
            hint["allyuziv-nom"] += 1
        elif label.startswith("iqtibos"):
            hint["iqtibos"] += 1
        entries.append({
            "id": f"{prefix}-{rid}",
            "type": type_,
            "unit": unit,
            "unit_normalized": normalize(unit),
            "pronunciations": json.dumps(parse_pronunciations(pron), ensure_ascii=False),
            "context_text": ctx_text,
            "intertext_author": it_author,
            "intertext_work": it_work,
            "intertext_genre": it_genre,
            "intertext_year": it_year,
            "intertext_publisher": it_pub,
            "intertext_page": it_page,
            "source_description": src_desc,
            "source_author": src_author,
            "source_work": src_work,
            "source_genre": src_genre,
            "source_publisher": src_pub,
            "source_period": src_period,
            "recognition": parse_recognition(recognition),
            "commentary": commentary,
            "semantic_field": parse_semantic_field(sem_field),
            "synonyms": json.dumps(parse_synonyms(synonym), ensure_ascii=False),
            "hypernym": hypernym,
            "hyponym": hyponym,
            "note": note,
            "is_complete": 1 if (src_desc or commentary or note) else 0,
        })
    return entries, hint


def read_rows_from_path(path: str, type_: str) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    try:
        return rows_from_workbook(wb, type_)[0]
    finally:
        wb.close()


def read_rows_from_bytes(data: bytes, type_: str) -> tuple[list[dict], dict]:
    wb = openpyxl.load_workbook(BytesIO(data))
    try:
        return rows_from_workbook(wb, type_)
    finally:
        wb.close()


def _insert(con: sqlite3.Connection, entries: list[dict]) -> None:
    cols = list(entries[0].keys())
    con.executemany(
        f"INSERT INTO entries ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
        [tuple(e[c] for c in cols) for e in entries],
    )


def rebuild(all_entries: list[dict]) -> None:
    """Bazani noldan qurish (CLI skript uchun)."""
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if DB_PATH.exists():
        DB_PATH.unlink()
    con = sqlite3.connect(DB_PATH)
    try:
        con.executescript(SCHEMA)
        _insert(con, all_entries)
        con.commit()
    finally:
        con.close()


def replace_type(entries: list[dict], type_: str) -> None:
    """Bitta tur yozuvlarini yangilash: eski shu turdagilar o'chib,
    yangilari yoziladi; boshqa tur yozuvlari saqlanadi (tranzaksiyada)."""
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    try:
        con.executescript(SCHEMA)
        con.execute("DELETE FROM entries WHERE type = ?", (type_,))
        _insert(con, entries)
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()
